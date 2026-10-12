import re
from docx import Document

d = Document(r"template\Thesis_Proposal_Form.docx")
t = d.tables[0]
out = open("tmp_inspect.txt", "w", encoding="utf-8")
for ri in (0, 1):
    xml = t.rows[ri].cells[1]._tc.xml
    xml = re.sub(r' xmlns:\w+="[^"]+"', "", xml)
    out.write(f"CELL{ri}\n{xml}\n")
trpr = re.sub(r' xmlns:\w+="[^"]+"', "", t.rows[0]._tr.xml)[:1500]
out.write("ROW0\n" + trpr + "\n")
out.write("TABLE STYLE XML\n")
for s in d.styles:
    if s.type == 3:
        out.write(s.name + "\n")
p = d.paragraphs[0]
out.write(re.sub(r' xmlns:\w+="[^"]+"', "", p._p.xml) + "\n")
h = d.paragraphs[4]
out.write(re.sub(r' xmlns:\w+="[^"]+"', "", h._p.xml) + "\n")
out.close()
