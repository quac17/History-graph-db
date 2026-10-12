# -*- coding: utf-8 -*-
"""Ghép các file Word từng mục hoặc từng chương thành một file lớn.

File đầu làm nền: khổ trang, lề, style. Các file sau chỉ chép phần thân.
Không đọc Markdown và không tạo lại nội dung.

Ví dụ:
  python "Do An/tools/ghep_docx.py" --doc thesis-proposal --lang vi
  python "Do An/tools/ghep_docx.py" --doc thesis --lang both
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TOOLS = Path(__file__).resolve().parent
DO_AN = TOOLS.parent
DOCS = {
    "thesis-proposal": DO_AN / "thesis-proposal",
    "thesis": DO_AN / "thesis",
}


def load_manifest(folder: Path) -> dict:
    path = folder / "manifest.json"
    if not path.is_file():
        raise FileNotFoundError("Không thấy %s" % path)
    return json.loads(path.read_text(encoding="utf-8"))


def insert_page_break(body, before):
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run.append(br)
    paragraph.append(run)
    if before is None:
        body.append(paragraph)
    else:
        before.addprevious(paragraph)


def next_rid(dst: Document) -> str:
    ids = []
    for rel in dst.part.rels.values():
        try:
            ids.append(int(rel.rId.replace("rId", "")))
        except ValueError:
            continue
    return "rId%d" % ((max(ids) if ids else 0) + 1)


def copy_rel(dst: Document, src: Document, old_rid: str) -> str:
    rels = src.part.rels
    if old_rid not in rels:
        return old_rid
    rel = rels[old_rid]
    new_rid = next_rid(dst)
    if rel.is_external:
        dst.part.rels.add_relationship(rel.reltype, rel.target_ref, new_rid, is_external=True)
        return new_rid
    dst.part.rels.add_relationship(rel.reltype, rel.target_part, new_rid)
    return new_rid


def remap_rids(element, dst: Document, src: Document, cache: dict):
    for node in element.iter():
        for attr in node.attrib:
            if attr.endswith("}embed") or attr.endswith("}id") or attr.endswith("}link"):
                old = node.get(attr)
                if not old or not old.startswith("rId"):
                    continue
                if old not in cache:
                    cache[old] = copy_rel(dst, src, old)
                node.set(attr, cache[old])


def append_part(dst: Document, src_path: Path, page_break: bool):
    src = Document(str(src_path))
    body = dst.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if page_break:
        insert_page_break(body, sect_pr)
    cache = {}
    for child in src.element.body:
        if child.tag == qn("w:sectPr"):
            continue
        copied = deepcopy(child)
        remap_rids(copied, dst, src, cache)
        if sect_pr is None:
            body.append(copied)
        else:
            sect_pr.addprevious(copied)


def merge(folder: Path, lang: str) -> Path:
    manifest = load_manifest(folder)
    if lang not in manifest:
        raise KeyError("manifest không có ngôn ngữ %s" % lang)
    spec = manifest[lang]
    parts = spec["parts"]
    missing = [folder / p["file"] for p in parts if not (folder / p["file"]).is_file()]
    if missing:
        names = "\n  ".join(str(p.name) for p in missing)
        raise FileNotFoundError("Thiếu file Word từng phần:\n  %s" % names)
    first = folder / parts[0]["file"]
    out = folder / spec["output"]
    tmp = folder / (spec["output"] + ".tmp.docx")
    shutil.copyfile(first, tmp)
    dst = Document(str(tmp))
    for part in parts[1:]:
        append_part(dst, folder / part["file"], bool(part.get("page_break")))
    dst.save(str(tmp))
    tmp.replace(out)
    return out


def self_test() -> None:
    folder = TOOLS / "_ghep_test"
    folder.mkdir(exist_ok=True)
    try:
        a = Document()
        a.add_paragraph("Phần A")
        a.save(str(folder / "a.docx"))
        b = Document()
        b.add_paragraph("Phần B")
        b.save(str(folder / "b.docx"))
        (folder / "manifest.json").write_text(
            json.dumps({
                "vi": {
                    "output": "full.docx",
                    "parts": [
                        {"file": "a.docx", "page_break": False},
                        {"file": "b.docx", "page_break": True},
                    ],
                }
            }),
            encoding="utf-8",
        )
        out = merge(folder, "vi")
        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        if texts != ["Phần A", "Phần B"]:
            raise RuntimeError("Ghép sai nội dung: %s" % texts)
        print("self-test ok", out)
    finally:
        shutil.rmtree(folder, ignore_errors=True)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Ghép Word từng phần thành file lớn")
    parser.add_argument("--doc", choices=sorted(DOCS), help="thesis hoặc thesis-proposal")
    parser.add_argument("--lang", choices=("vi", "en", "both"), default="vi")
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args(argv)
    if args.self_test:
        self_test()
        return 0
    if not args.doc:
        parser.error("cần --doc hoặc --self-test")
    folder = DOCS[args.doc]
    langs = ("vi", "en") if args.lang == "both" else (args.lang,)
    for lang in langs:
        path = merge(folder, lang)
        print(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
