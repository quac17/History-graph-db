# -*- coding: utf-8 -*-
"""Dựng slide proposal từ khung markdown. Slide đơn giản, không ảnh minh họa."""
from __future__ import annotations

import math
import os
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
KHUNG = ROOT / "khung-slide-bao-cao-proposal.md"
OUT = ROOT / "slide-bao-cao-proposal.pptx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5B, 0x64, 0x70)
ROW = RGBColor(0xF4, 0xF7, 0xFB)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
LEFT = Inches(0.48)
WIDTH = Inches(12.37)
TOP = Inches(1.12)
BOTTOM = Inches(6.82)


def set_run(run, text, size, bold=False, italic=False, color=INK):
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = r_pr.makeelement(qn(tag))
            r_pr.append(node)
        node.set("typeface", "Arial")


def add_box(slide, left, top, width, height):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    return tf


def write_lines(tf, items, size, color=INK, bold=False, bullet=True, space=6):
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = item.get("level", 0)
        p.space_after = Pt(space)
        p.space_before = Pt(0)
        prefix = ""
        if bullet:
            marker = item.get("marker", "-")
            if marker == "-":
                prefix = "• " if item.get("level", 0) == 0 else "– "
            else:
                prefix = marker + " "
        set_run(p.add_run(), prefix + item["text"], size, bold=bold, color=color)


def clean(text):
    return text.replace("`", "")


def is_item(line):
    return bool(re.match(r"^\s*(?:-|\d+\.)\s+", line))


def parse_item(line):
    m = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", line)
    level = 1 if len(m.group(1)) >= 2 else 0
    return {"level": level, "marker": m.group(2), "text": clean(m.group(3).strip())}


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [clean(c.strip()) for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows


def parse_khung(text):
    slides = []
    section = ""
    current = None
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# Phần "):
            section = line[2:].strip()
            i += 1
            continue
        if line.startswith("## Slide "):
            title = re.sub(r"^## Slide \d+\.\s*", "", line).strip()
            num = int(re.search(r"Slide (\d+)", line).group(1))
            current = {"num": num, "title": title, "section": section, "blocks": [], "note": ""}
            slides.append(current)
            i += 1
            continue
        if line.startswith("# ") or (line.startswith("## ") and not line.startswith("## Slide ")):
            current = None
            i += 1
            continue
        if current is None:
            i += 1
            continue
        if not line.strip() or line.strip() in ("---", "***") or line.strip() == "Nội dung:" or line.startswith("Gợi ý trình bày"):
            i += 1
            continue
        if line.startswith("Lời nói:"):
            current["note"] = line.split(":", 1)[1].strip()
            i += 1
            continue
        if line.startswith("|"):
            chunk = []
            while i < len(lines) and lines[i].startswith("|"):
                chunk.append(lines[i])
                i += 1
            current["blocks"].append(("table", parse_table(chunk)))
            continue
        if is_item(line):
            items = []
            while i < len(lines) and is_item(lines[i]):
                items.append(parse_item(lines[i]))
                i += 1
            current["blocks"].append(("list", items))
            continue
        text_line = line.strip()
        j = i + 1
        while j < len(lines) and not lines[j].strip():
            j += 1
        if text_line.endswith(":") and j < len(lines) and is_item(lines[j]):
            items = []
            i = j
            while i < len(lines) and is_item(lines[i]):
                items.append(parse_item(lines[i]))
                i += 1
            current["blocks"].append(("group", text_line, items))
            continue
        current["blocks"].append(("text", clean(text_line)))
        i += 1
    return slides


def chrome(slide, spec, total):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    if spec["section"]:
        tf = add_box(slide, Inches(7.2), Inches(0.16), Inches(5.65), Inches(0.32))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        set_run(p.add_run(), spec["section"].replace(". ", " · "), 12, color=NAVY)
    tf = add_box(slide, LEFT, Inches(0.42), WIDTH, Inches(0.62))
    p = tf.paragraphs[0]
    set_run(p.add_run(), spec["title"], 26, bold=True, color=NAVY)
    foot = add_box(slide, LEFT, Inches(7.05), Inches(9.2), Inches(0.28))
    p = foot.paragraphs[0]
    set_run(p.add_run(), "Trần Việt Hoàng · Proposal đồ án", 11, color=MUTED)
    num = add_box(slide, Inches(11.4), Inches(7.05), Inches(1.45), Inches(0.28))
    p = num.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.add_run(), "%d / %d" % (spec["num"], total), 11, color=MUTED)
    if spec["note"]:
        notes = slide.notes_slide.notes_text_frame
        notes.text = spec["note"]


def cover(prs, spec, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SLIDE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    items = []
    for kind, payload in spec["blocks"]:
        if kind == "list":
            items.extend(payload)
    lines = [it["text"] for it in items]
    school, title_vi, title_en, student, advisor, place = (lines + [""] * 6)[:6]
    tf = add_box(slide, Inches(0.7), Inches(1.15), Inches(11.8), Inches(0.4))
    set_run(tf.paragraphs[0].add_run(), school, 16, color=NAVY)
    tf = add_box(slide, Inches(0.7), Inches(1.85), Inches(11.8), Inches(1.7))
    set_run(tf.paragraphs[0].add_run(), title_vi, 30, bold=True, color=NAVY)
    tf = add_box(slide, Inches(0.7), Inches(3.7), Inches(11.8), Inches(1.05))
    set_run(tf.paragraphs[0].add_run(), title_en, 16, italic=True, color=MUTED)
    tf = add_box(slide, Inches(0.7), Inches(5.15), Inches(11.8), Inches(1.15))
    p = tf.paragraphs[0]
    set_run(p.add_run(), student, 16, color=INK)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    set_run(p2.add_run(), advisor, 16, color=INK)
    tf = add_box(slide, Inches(0.7), Inches(6.55), Inches(11.8), Inches(0.32))
    set_run(tf.paragraphs[0].add_run(), place, 14, color=MUTED)
    foot = add_box(slide, Inches(0.7), Inches(7.05), Inches(9.2), Inches(0.28))
    set_run(foot.paragraphs[0].add_run(), "Trần Việt Hoàng · Proposal đồ án", 11, color=MUTED)
    num = add_box(slide, Inches(11.4), Inches(7.05), Inches(1.45), Inches(0.28))
    p = num.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.add_run(), "1 / %d" % total, 11, color=MUTED)


def style_cell(cell, text, size, bold, fill, font_color, center=False):
    cell.text = ""
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    set_run(p.add_run(), text, size, bold=bold, color=font_color)


def column_widths(rows, width):
    cols = len(rows[0])
    weights = []
    for c in range(cols):
        longest = max(len(row[c]) for row in rows)
        weights.append(max(longest, 10))
    total = float(sum(weights))
    widths = [width * w / total for w in weights]
    floor = width * 0.12
    if any(w < floor for w in widths):
        widths = [max(w, floor) for w in widths]
        scale = width / sum(widths)
        widths = [w * scale for w in widths]
    if cols >= 3:
        last_len = max(len(row[-1]) for row in rows)
        need = int(Inches(min(3.05, max(1.6, last_len * 0.092))))
        if widths[-1] < need and sum(widths[:-1]) > need:
            rest = width - need
            head = sum(widths[:-1])
            widths = [int(w / head * rest) for w in widths[:-1]] + [need]
    return widths


def row_line_count(text, col_width):
    chars = max(10, col_width / 914400 * 72 / 6.4)
    return max(1, math.ceil(len(text) / chars))


def add_table(slide, rows, left, top, width, height):
    widths = column_widths(rows, width)
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, height)
    table = table_shape.table
    highlight = rows[0][-1] == "Đồ án"
    for c, col_w in enumerate(widths):
        table.columns[c].width = int(col_w)
    for r, row in enumerate(rows):
        lines = 1
        for c, text in enumerate(row):
            lines = max(lines, row_line_count(text, widths[c]))
        table.rows[r].height = Inches(0.18 + lines * 0.24)
        for c, text in enumerate(row):
            last = highlight and c == len(row) - 1
            if r == 0:
                style_cell(table.cell(r, c), text, 12, True, NAVY, WHITE, center=True)
            else:
                fill = RGBColor(0xE4, 0xEC, 0xF6) if last else (WHITE if r % 2 else ROW)
                style_cell(table.cell(r, c), text, 12, last, fill, INK, center=False)
    return table_shape


def table_height(rows, width):
    widths = column_widths(rows, width)
    total = 0.0
    for row in rows:
        lines = 1
        for text, col_w in zip(row, widths):
            lines = max(lines, row_line_count(text, col_w))
        total += 0.18 + lines * 0.24
    return total


def text_height(items, width_in, size):
    chars = max(18, width_in * 72 / (size * 0.52))
    lines = 0
    for item in items:
        use = chars * (0.86 if item.get("level") else 1)
        lines += max(1, math.ceil(len(item["text"]) / use))
    return lines * (size * 1.35 / 72) + 0.08 * max(0, len(items) - 1)


def render_list(slide, items, left, top, width, height, size=16):
    tf = add_box(slide, left, top, width, height)
    write_lines(tf, items, size, space=8 if size >= 16 else 4)


def render_group(slide, title, items, left, top, width, height, size=16):
    box_h = Inches(0.34)
    head = add_box(slide, left, top, width, box_h)
    set_run(head.paragraphs[0].add_run(), title, size, bold=True, color=NAVY)
    render_list(slide, items, left, top + box_h, width, height - box_h, size)


def stack(slide, blocks):
    gap = 0.12
    sizes = []
    estimates = []
    for block in blocks:
        kind = block[0]
        if kind == "table":
            estimates.append(table_height(block[1], WIDTH))
            sizes.append(12)
        elif kind == "list":
            size = 18 if len(block[1]) <= 4 else 16
            if len(block[1]) >= 7:
                size = 14
            estimates.append(text_height(block[1], 12.2, size) + 0.08)
            sizes.append(size)
        elif kind == "group":
            size = 16 if len(block[2]) <= 5 else 14
            estimates.append(0.38 + text_height(block[2], 12.2, size))
            sizes.append(size)
        else:
            estimates.append(0.4)
            sizes.append(15)
    avail = (BOTTOM - TOP) / 914400
    total = sum(estimates) + gap * max(0, len(blocks) - 1)
    overflow = max(0, total - avail)
    if overflow and any(b[0] != "table" for b in blocks):
        flexible = sum(est for block, est in zip(blocks, estimates) if block[0] != "table")
        if flexible:
            keep = max(0.7, (flexible - overflow) / flexible)
            estimates = [est if block[0] == "table" else est * keep for block, est in zip(blocks, estimates)]
    y = TOP
    for block, est, size in zip(blocks, estimates, sizes):
        h = Inches(est)
        if block[0] == "table":
            add_table(slide, block[1], LEFT, y, WIDTH, h)
        elif block[0] == "list":
            render_list(slide, block[1], LEFT, y, WIDTH, h, size)
        elif block[0] == "group":
            render_group(slide, block[1], block[2], LEFT, y, WIDTH, h, size)
        else:
            tf = add_box(slide, LEFT, y, WIDTH, h)
            set_run(tf.paragraphs[0].add_run(), block[1], size, color=INK)
        y += h + Inches(gap)


def two_columns(slide, blocks):
    col_w = Inches(5.95)
    gap = Inches(0.47)
    height = BOTTOM - TOP
    for i, block in enumerate(blocks):
        left = LEFT + i * (col_w + gap)
        if block[0] == "group":
            render_group(slide, block[1], block[2], left, TOP, col_w, height, 16)
        elif block[0] == "list":
            render_list(slide, block[1], left, TOP, col_w, height, 16)
        elif block[0] == "table":
            add_table(slide, block[1], left, TOP, col_w, height)


def build(spec, prs, total):
    if spec["num"] == 1:
        cover(prs, spec, total)
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(slide, spec, total)
    blocks = spec["blocks"]
    groups = [b for b in blocks if b[0] in ("group", "list")]
    tables = [b for b in blocks if b[0] == "table"]
    if len(blocks) == 2 and len(groups) == 2 and not tables:
        two_columns(slide, blocks)
    elif spec["num"] == 17 and tables and groups:
        table_w = Inches(7.45)
        height = Inches(table_height(tables[0][1], table_w))
        add_table(slide, tables[0][1], LEFT, TOP, table_w, height)
        render_group(slide, groups[0][1], groups[0][2], Inches(8.15), TOP, Inches(4.7), Inches(5.4), 14)
    else:
        stack(slide, blocks)


def main():
    slides = parse_khung(KHUNG.read_text(encoding="utf-8"))
    if len(slides) != 20:
        raise SystemExit("Khung có %d slide, cần 20" % len(slides))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "Thesis Proposal — Trần Việt Hoàng"
    prs.core_properties.author = "Trần Việt Hoàng"
    for spec in slides:
        build(spec, prs, len(slides))
    tmp = OUT.with_suffix(".tmp.pptx")
    prs.save(str(tmp))
    os.replace(tmp, OUT)
    print(OUT)
    print("slides", len(slides))


if __name__ == "__main__":
    sys.exit(main())
