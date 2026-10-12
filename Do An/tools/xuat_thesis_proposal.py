# -*- coding: utf-8 -*-
"""Xuất từng mục thesis proposal từ Markdown sang Word.

Đọc muc-0N.md thành muc-0N.vi.docx và muc-0N.en.md thành muc-0N.en.docx.
Không ghép file lớn. Ghép bằng ghep_docx.py.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

TOOLS = Path(__file__).resolve().parent
FOLDER = TOOLS.parent / "thesis-proposal"
BLACK = RGBColor(0, 0, 0)
TABLE_W = 9773
TIMELINE_WIDTHS = (1410, 1984, 3544, 2835)
FONT = "Times New Roman"


def set_run_font(run, size, bold=False, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), FONT)


def paragraph_format(p, before=5, after=5, line=1.0, align="justify", left=0, hanging=0):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.left_indent = Cm(left) if left else Cm(0)
    fmt.first_line_indent = Cm(-hanging) if hanging else Cm(0)


def add_text(p, text, size=12, bold=False, italic=False):
    run = p.add_run(text)
    set_run_font(run, size, bold, italic)
    return run


def plain(text):
    return text.replace("`", "")


def parse_blocks(text):
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
            i += 1
            continue
        if line.startswith("# "):
            blocks.append(("title", line[2:].strip()))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [plain(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(plain(lines[i][2:].strip()))
                i += 1
            blocks.append(("bullets", items))
            continue
        stripped = line.strip()
        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            blocks.append(("italic", plain(stripped[1:-1].strip())))
            i += 1
            continue
        blocks.append(("p", plain(stripped)))
        i += 1
    return blocks


def shade_header_off(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    v_align = OxmlElement("w:vAlign")
    v_align.set(qn("w:val"), "center")
    tc_pr.append(v_align)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "C4C7C5")
        borders.append(el)
    tc_pr.append(borders)


def set_table_fixed(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_W))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", "119"), ("left", "181"), ("bottom", "119"), ("right", "181")):
        node = OxmlElement("w:%s" % name)
        node.set(qn("w:w"), val)
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)
    grid = tbl.find(qn("w:tblGrid"))
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr = row._tr
        tr_pr = tr.get_or_add_trPr()
        for old in tr_pr.findall(qn("w:tblHeader")):
            tr_pr.remove(old)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_pr.append(header_flag)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            shade_header_off(cell)
            tc = cell._tc
            tc_pr = tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def column_widths(headers):
    n = len(headers)
    if n == 4 and headers[0] in ("Tuần", "Week"):
        return TIMELINE_WIDTHS
    if n == 3:
        return (2500, 4873, 2400)
    share = TABLE_W // n
    widths = [share] * n
    widths[-1] = TABLE_W - share * (n - 1)
    return tuple(widths)


def fill_cell(cell, text, bold, center):
    cell.text = ""
    p = cell.paragraphs[0]
    paragraph_format(p, before=0, after=0, align="center" if center else "left")
    add_text(p, text, 12, bold=bold)


def add_table(doc, rows):
    headers = rows[0]
    widths = column_widths(headers)
    timeline = headers[0] in ("Tuần", "Week")
    table = doc.add_table(rows=len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_fixed(table, widths)
    for r_i, row in enumerate(rows):
        for c_i, text in enumerate(row):
            bold = r_i == 0 or (timeline and c_i < 2)
            fill_cell(table.rows[r_i].cells[c_i], text, bold, center=(r_i == 0))
    return table


def add_label_or_body(doc, text):
    match = re.match(r"^\*\*(.+?)\*\*\s*(.*)$", text)
    p = doc.add_paragraph()
    paragraph_format(p)
    if not match:
        add_text(p, text, 12)
        return
    label = match.group(1).strip()
    if not label.endswith(":"):
        label += ":"
    add_text(p, label, 12, bold=True)
    rest = match.group(2).strip()
    if rest:
        add_text(p, " " + rest, 12)


def render(md_path, out_path):
    blocks = parse_blocks(md_path.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(1.27)
    section.footer_distance = Cm(1.27)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    r_pr = normal.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), FONT)
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)
    for kind, payload in blocks:
        if kind == "title":
            p = doc.add_paragraph()
            paragraph_format(p, before=0, after=0, line=1.15, align="center")
            add_text(p, payload, 18, bold=True)
        elif kind == "h2":
            p = doc.add_paragraph()
            paragraph_format(p, align="left")
            add_text(p, payload, 12, bold=True)
        elif kind == "italic":
            p = doc.add_paragraph()
            paragraph_format(p)
            add_text(p, payload, 12, italic=True)
        elif kind == "bullets":
            for item in payload:
                p = doc.add_paragraph()
                paragraph_format(p, left=1.27, hanging=0.63)
                add_text(p, "• " + item, 12)
        elif kind == "table":
            add_table(doc, payload)
        else:
            add_label_or_body(doc, payload)
    doc.save(str(out_path))


def main():
    written = []
    for n in range(1, 7):
        vi = FOLDER / ("muc-%02d.md" % n)
        en = FOLDER / ("muc-%02d.en.md" % n)
        vi_out = FOLDER / ("muc-%02d.vi.docx" % n)
        en_out = FOLDER / ("muc-%02d.en.docx" % n)
        render(vi, vi_out)
        render(en, en_out)
        written.extend((vi_out, en_out))
    for path in written:
        print(path.name)
    return 0


if __name__ == "__main__":
    sys.exit(main())
